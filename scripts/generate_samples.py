from pathlib import Path
import csv, json, shutil
from openpyxl import Workbook
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image

root = Path(__file__).resolve().parents[1]
out = root / "apps/web/public/samples"
out.mkdir(parents=True, exist_ok=True)
headers = [
    "Họ và tên",
    "Năm sinh",
    "Đơn vị hiện tại",
    "Cấp bậc",
    "Số hiệu",
    "Hưởng lương",
    "Nhóm đối tượng",
    "Trạng thái công tác",
    "Số tháng đóng góp",
    "Chế độ yêu cầu",
    "Nguồn dữ liệu",
]
rows = [
    [
        "Nguyễn Minh An",
        "1985",
        "Cục Cảnh sát giao thông",
        "Đại úy",
        "DEMO-001",
        "Có",
        "BCA_SI_QUAN",
        "ACTIVE",
        "240",
        "DEMO_RETIREMENT",
        "SYNTHETIC_DEMO",
    ],
    [
        "Trần Hải Bình",
        "1987",
        "Bộ Tổng Tham mưu",
        "Thiếu tá",
        "DEMO-002",
        "Có",
        "BQP_SI_QUAN",
        "ACTIVE",
        "180",
        "DEMO_RETIREMENT",
        "SYNTHETIC_DEMO",
    ],
]
with (out / "danh_sach.csv").open("w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(headers)
    w.writerows(rows)
w = Workbook()
sheet = w.active
sheet.title = "Synthetic demo"
sheet.append(headers)
for row in rows:
    sheet.append(row)
for col in sheet.columns:
    sheet.column_dimensions[col[0].column_letter].width = 28
w.save(out / "danh_sach.xlsx")
d = Document()
d.add_heading("HỒ SƠ MẪU — DỮ LIỆU MÔ PHỎNG", 0)
for row in rows:
    for k, v in zip(headers, row):
        d.add_paragraph(f"{k}: {v}")
    d.add_paragraph("")
d.save(out / "ho_so.docx")
pdfmetrics.registerFont(
    TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
)
c = canvas.Canvas(str(out / "ho_so_text.pdf"))
c.setTitle("Synthetic sample dossier")
for row in rows:
    c.setFont("DejaVu", 12)
    c.drawString(40, 800, "HỒ SƠ MẪU — DỮ LIỆU MÔ PHỎNG")
    y = 755
    for k, v in zip(headers, row):
        c.drawString(40, y, f"{k}: {v}")
        y -= 30
    c.showPage()
c.save()
images = []
for name in ["bca_sample_image.jpg", "bqp_sample_image.jpg"]:
    shutil.copyfile(root.parent / "image_cases" / name, out / name)
    images.append(Image.open(out / name).convert("RGB"))
scan = canvas.Canvas(str(out / "ho_so_scan.pdf"))
for name, im in zip(["bca_sample_image.jpg", "bqp_sample_image.jpg"], images):
    scan.setPageSize(im.size)
    scan.drawImage(str(out / name), 0, 0, width=im.width, height=im.height)
    scan.showPage()
scan.save()
entries = [
    ("csv", "CSV", "danh_sach.csv", "text/csv"),
    (
        "xlsx",
        "XLSX",
        "danh_sach.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ),
    (
        "docx",
        "DOCX",
        "ho_so.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
    ("pdf-text", "PDF text", "ho_so_text.pdf", "application/pdf"),
    ("pdf-scan", "PDF scan · 2 trang", "ho_so_scan.pdf", "application/pdf"),
    ("image-bca", "Ảnh BCA", "bca_sample_image.jpg", "image/jpeg"),
    ("image-bqp", "Ảnh BQP", "bqp_sample_image.jpg", "image/jpeg"),
]
manifest = {
    "notice": "Generated identities and policy facts are SYNTHETIC_DEMO. Supplied images are PROVIDED, OCR may require review. Unit references do not verify any individual.",
    "as_of_date": "2026-09-16",
    "sources": [
        "https://www.bocongan.gov.vn/bai-viet/cuc-canh-sat-giao-thong-thong-tin-tien-do-lam-sach-du-lieu-cac-loai-xe-may-bien-so-cu-3-so-4-so-1759747160",
        "https://www.mod.gov.vn/vn/chi-tiet/sa-ttsk/sa-tt-qpan/bo-tong-tham-muu-quan-doi-nhan-dan-viet-nam-voi-su-nghiep-dau-tranh-giai-phong-dan-toc-xay-dung-va-bao-ve-to-quoc",
    ],
    "files": [
        dict(
            id=i,
            label=l,
            filename=f,
            content_type=m,
            source_kind="PROVIDED"
            if i.startswith("image") or i == "pdf-scan"
            else "SYNTHETIC_DEMO",
        )
        for i, l, f, m in entries
    ],
    "manual": [
        dict(
            zip(
                [
                    "fullName",
                    "birthYear",
                    "department",
                    "position",
                    "identifier",
                    "salary",
                    "subjectGroup",
                    "employmentStatus",
                    "contributionMonths",
                    "requestedRegime",
                    "sourceKind",
                ],
                row,
            ),
            asOfDate="2026-09-16",
            extraInfo="Hồ sơ mẫu SYNTHETIC_DEMO; hưởng lương",
        )
        for row in rows
    ],
}
(out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
