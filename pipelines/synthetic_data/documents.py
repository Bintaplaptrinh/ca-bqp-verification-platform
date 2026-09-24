from __future__ import annotations

import csv
from pathlib import Path

import pandas as pd
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


_FONT_CANDIDATES = (
    Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
    Path('/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf'),
    Path('/usr/share/fonts/opentype/noto/NotoSans-Regular.ttf'),
)


def _unicode_font_path() -> Path:
    for path in _FONT_CANDIDATES:
        if path.exists():
            return path
    raise RuntimeError(
        'Unicode font is required to generate Vietnamese OCR fixtures. '
        'Install fonts-dejavu-core or Noto Sans instead of stripping accents.'
    )


def _pdf_font_name() -> str:
    name = 'SyntheticVietnameseUnicode'
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(name, str(_unicode_font_path())))
    return name


def _pil_font(size: int = 28):
    return ImageFont.truetype(str(_unicode_font_path()), size=size)


def _draw_wrapped(draw: ImageDraw.ImageDraw, text: str, xy: tuple[int, int], *, font, max_width: int, line_gap: int = 10):
    x, y = xy
    words = text.split()
    line = ''
    for word in words:
        candidate = f'{line} {word}'.strip()
        box = draw.textbbox((0, 0), candidate, font=font)
        if line and box[2] - box[0] > max_width:
            draw.text((x, y), line, fill=35, font=font)
            y += (draw.textbbox((0, 0), line, font=font)[3] + line_gap)
            line = word
        else:
            line = candidate
    if line:
        draw.text((x, y), line, fill=35, font=font)


def generate(records_csv: Path, out_dir: Path, limit: int, batch_size: int):
    del batch_size  # retained for CLI/API compatibility
    df = pd.read_csv(records_csv, dtype=str).fillna('').head(limit)
    out_dir.mkdir(parents=True, exist_ok=True)

    for sub in ['txt', 'docx', 'xlsx', 'pdf_text', 'png_scan']:
        (out_dir / sub).mkdir(parents=True, exist_ok=True)

    manifest = []
    pdf_font = _pdf_font_name()
    pil_font = _pil_font()

    for _, row in df.iterrows():
        rid = row['record_id']
        text = row['raw_text']

        txt = out_dir / 'txt' / f'{rid}.txt'
        txt.write_text(text, encoding='utf-8')

        docx = out_dir / 'docx' / f'{rid}.docx'
        doc = Document()
        doc.add_heading('HỒ SƠ SYNTHETIC DEMO', 1)
        doc.add_paragraph(text)
        doc.save(docx)

        xlsx = out_dir / 'xlsx' / f'{rid}.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.append(['record_id', 'full_name', 'personal_code', 'raw_text'])
        ws.append([rid, row['full_name'], row['personal_code'], text])
        wb.save(xlsx)

        pdf = out_dir / 'pdf_text' / f'{rid}.pdf'
        c = canvas.Canvas(str(pdf))
        c.setFont(pdf_font, 10)
        text_obj = c.beginText(50, 800)
        text_obj.setFont(pdf_font, 10)
        for offset in range(0, len(text), 110):
            text_obj.textLine(text[offset:offset + 110])
        c.drawText(text_obj)
        c.save()

        png = out_dir / 'png_scan' / f'{rid}.png'
        img = Image.new('L', (1200, 700), 245)
        draw = ImageDraw.Draw(img)
        _draw_wrapped(draw, text, (70, 80), font=pil_font, max_width=1060)
        img.save(png)

        for kind, path in [('txt', txt), ('docx', docx), ('xlsx', xlsx), ('pdf_text', pdf), ('png_scan', png)]:
            manifest.append({'record_id': rid, 'document_type': kind, 'path': str(path.relative_to(out_dir))})

    with (out_dir / 'document_manifest.csv').open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=['record_id', 'document_type', 'path'])
        w.writeheader()
        w.writerows(manifest)
