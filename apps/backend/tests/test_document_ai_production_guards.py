from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import fitz
import pytest
from openpyxl import Workbook
from PIL import Image

from cabqp.modules.bulk import service as bulk_service
from cabqp.modules.document_intelligence import parsers, router
from cabqp.modules.document_intelligence.extraction import extract
from cabqp.modules.document_intelligence.limits import DocumentLimitError
from cabqp.modules.document_intelligence.router import RoutingDecision
from cabqp.shared.enums import InputKind


def test_single_high_confidence_unit_cannot_make_document_confidence_near_one():
    result = extract("Đơn vị công tác: Cục Kỹ thuật")

    assert result.current_unit == "Cục Kỹ thuật"
    assert result.fields["field_confidence"]["current_unit"] > 0.8
    assert result.extraction_confidence < 0.75


def test_parser_tables_feed_generic_field_extraction_without_value_repair():
    tables = [[
        ["Họ và tên", "Đỗ Minh Quân"],
        ["CCCD", "075201223344"],
        ["Chức vụ", "Chuyên viên"],
        ["Đơn vị công tác", "Cục Kỹ thuật"],
    ]]
    result = extract("", {"_parse_method": "DOCX_TEXT", "_parsed_tables": tables})

    assert result.subject_name == "Đỗ Minh Quân"
    assert result.subject_code == "075201223344"
    assert result.position == "Chuyên viên"
    assert result.current_unit == "Cục Kỹ thuật"
    assert result.extraction_confidence > 0.85
    assert result.fields["field_evidence"]["subject_name"]["rule"] == "table_right_cell"


def _multi_sheet_xlsx() -> bytes:
    wb = Workbook()
    instructions = wb.active
    instructions.title = "Instructions"
    instructions.append(["Read me"])
    data = wb.create_sheet("Data")
    data.append(["Họ và tên", "CCCD", "Chức vụ", "Đơn vị công tác"])
    data.append(["Nguyễn Văn A", "001082946357", "Chuyên viên", "Cục Kỹ thuật"])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_bulk_reader_ingests_data_from_non_active_worksheet():
    headers, rows, _ = bulk_service._read_rows("people.xlsx", _multi_sheet_xlsx())

    assert "Họ và tên" in headers
    assert len(rows) == 1
    assert rows[0]["Họ và tên"] == "Nguyễn Văn A"
    assert rows[0]["_source_sheet"] == "Data"
    assert rows[0]["_source_row"] == 2


def test_key_value_spreadsheet_keeps_rows_for_source_preview():
    result=parsers.parse_document("people.xlsx",_multi_sheet_xlsx())

    preview=result.evidence["source_preview_rows"]
    assert ["Họ và tên","CCCD","Chức vụ","Đơn vị công tác"] in preview
    assert ["Nguyễn Văn A","001082946357","Chuyên viên","Cục Kỹ thuật"] in preview


def test_plain_text_keeps_content_for_source_preview():
    text="Họ và tên: Nguyễn Văn Minh\nĐơn vị: Cục Kỹ thuật"
    result=parsers.parse_document("person.txt",text.encode("utf-8"))

    assert result.evidence["source_preview_text"] == text


def test_pdf_page_limit_is_checked_before_per_page_routing(monkeypatch):
    doc = fitz.open()
    doc.new_page()
    doc.new_page()
    content = doc.tobytes()
    doc.close()
    monkeypatch.setattr(router, "get_settings", lambda: SimpleNamespace(document_max_pdf_pages=1))

    with pytest.raises(DocumentLimitError) as exc:
        router.probe_pdf(content)
    assert exc.value.code == "PDF_PAGE_LIMIT_EXCEEDED"


def test_image_pixel_limit_rejects_before_ocr_model_runs(monkeypatch):
    out = BytesIO()
    Image.new("RGB", (20, 20), "white").save(out, format="PNG")
    settings = SimpleNamespace(
        document_soft_timeout_seconds=60,
        document_max_image_pixels=100,
    )
    monkeypatch.setattr(parsers, "get_settings", lambda: settings)
    monkeypatch.setattr(
        parsers,
        "route_input",
        lambda *_: RoutingDecision(InputKind.IMAGE, "png", "test"),
    )

    with pytest.raises(DocumentLimitError) as exc:
        parsers.parse_document("large.png", out.getvalue())
    assert exc.value.code == "IMAGE_PIXEL_LIMIT_EXCEEDED"


def test_hybrid_failed_page_caps_document_confidence(monkeypatch):
    from cabqp.modules.document_intelligence.ocr.base import OcrLine
    from cabqp.modules.document_intelligence.ocr.pipeline import OcrRun
    from cabqp.modules.document_intelligence.router import PageProbe

    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "Digital text " * 80)
    doc.new_page()
    content = doc.tobytes()
    doc.close()

    decision = RoutingDecision(
        InputKind.PDF_HYBRID,
        "pdf",
        "test",
        page_route_map=[
            PageProbe(1, "TEXT", 100, 0.0, 1, 4.0, 0.0, 0, []),
            PageProbe(2, "OCR", 0, 1.0, 0, 0.0, 0.0, 0, []),
        ],
    )
    settings = SimpleNamespace(
        document_soft_timeout_seconds=60,
        document_max_ocr_pages=5,
        document_max_image_pixels=40_000_000,
        document_max_raster_pixels=100_000_000,
        scan_table_detection_enabled=False,
        document_failed_gate_confidence_cap=0.5,
        document_max_tables=20,
    )
    monkeypatch.setattr(parsers, "route_input", lambda *_: decision)
    monkeypatch.setattr(parsers, "get_settings", lambda: settings)
    monkeypatch.setattr(parsers, "confidence_from_quality", lambda *_: 0.95)
    monkeypatch.setattr(
        parsers,
        "run_ocr",
        lambda *_, **__: OcrRun(
            [OcrLine("Họ và tên: Nguyễn Văn A", 0.99, (0, 0, 100, 20), 2, "fake")],
            {"gate_result": "FAIL", "metrics": {}, "reason": "TEST_FAIL"},
            0.99,
            {"selected_engine": "fake", "critical_disagreement": False, "fallback_ran": False},
        ),
    )
    monkeypatch.setattr(parsers, "extract_pdf_table_results", lambda *_: [])

    result = parsers.parse_document("hybrid.pdf", content)

    assert result.quality["gate_result"] == "FAIL"
    assert result.confidence <= settings.document_failed_gate_confidence_cap


def test_record_table_detection_keeps_people_as_independent_rows():
    table = [
        ["STT", "Họ và tên", "CCCD", "Chức vụ", "Đơn vị công tác"],
        [1, "Nguyễn Văn A", "001082946357", "Chuyên viên", "Cục Kỹ thuật"],
        [2, "Trần Thị B", "079203001122", "Kỹ sư", "Trung tâm A"],
    ]

    assert bulk_service.tables_have_record_list([table]) is True


def test_vertical_key_value_form_is_not_misclassified_as_people_list():
    table = [
        ["Họ và tên", "Nguyễn Văn A"],
        ["CCCD", "001082946357"],
        ["Chức vụ", "Chuyên viên"],
        ["Đơn vị công tác", "Cục Kỹ thuật"],
    ]

    assert bulk_service.tables_have_record_list([table]) is False


def test_pdf_bulk_reader_uses_each_table_row_as_one_record(monkeypatch):
    table = [
        ["STT", "Họ và tên", "CCCD", "Chức vụ", "Đơn vị công tác"],
        [1, "Nguyễn Văn A", "001082946357", "Chuyên viên", "Cục Kỹ thuật"],
        [2, "Trần Thị B", "079203001122", "Kỹ sư", "Trung tâm A"],
    ]
    fake = SimpleNamespace(
        tables=[table],
        evidence={"table_sources": [{"source": "pdf-native", "page": 1}]},
    )
    monkeypatch.setattr(parsers, "parse_document", lambda *_: fake)

    headers, rows, _ = bulk_service._read_rows("people.pdf", b"fake")

    assert "Họ và tên" in headers
    assert [row["Họ và tên"] for row in rows] == ["Nguyễn Văn A", "Trần Thị B"]
    assert all(row["_source_table"] == 1 for row in rows)
    assert all(row["_source_page"] == 1 for row in rows)


def test_pdf_continuation_table_without_repeated_header_reuses_previous_schema(monkeypatch):
    first = [
        ["STT", "Họ và tên", "CCCD"],
        [1, "Nguyễn Văn A", "001082946357"],
        [2, "Trần Thị B", "079203001122"],
    ]
    continuation = [
        [3, "Lê Văn C", "012345678901"],
        [4, "Phạm Thị D", "012345678902"],
    ]
    fake = SimpleNamespace(
        tables=[first, continuation],
        evidence={"table_sources": [{"source": "pdf-native", "page": 1}, {"source": "pdf-native", "page": 2}]},
    )
    monkeypatch.setattr(parsers, "parse_document", lambda *_: fake)

    _, rows, _ = bulk_service._read_rows("people.pdf", b"fake")

    assert [row["Họ và tên"] for row in rows] == ["Nguyễn Văn A", "Trần Thị B", "Lê Văn C", "Phạm Thị D"]
    assert rows[-1]["_source_page"] == 2


def test_pp_structure_html_is_reconstructed_as_grid_with_confidence():
    from cabqp.modules.document_intelligence.table import tables_from_scan_structure

    structure = [{
        "res": {
            "table_res_list": [{
                "table_id": 0,
                "pred_html": "<table><tr><th>STT</th><th>Họ và tên</th></tr><tr><td>1</td><td>Nguyễn Văn A</td></tr><tr><td>2</td><td>Trần Thị B</td></tr></table>",
                "table_ocr_pred": {"rec_scores": [0.98, 0.97, 0.95, 0.94, 0.92, 0.91]},
            }]
        }
    }]

    tables = tables_from_scan_structure(structure)

    assert tables[0]["cells"][0] == ["STT", "Họ và tên"]
    assert tables[0]["cells"][2][1] == "Trần Thị B"
    assert 0.0 < tables[0]["confidence"] <= 1.0


def test_one_person_horizontal_roster_is_still_a_record_list():
    table = [
        ["Họ và tên", "CCCD", "Đơn vị công tác"],
        ["Nguyễn Văn A", "001082946357", "Cục Kỹ thuật"],
    ]

    assert bulk_service.tables_have_record_list([table]) is True


def test_empty_pp_structure_table_list_is_not_a_table():
    from cabqp.modules.document_intelligence.table import scan_structure_contains_table

    assert scan_structure_contains_table([{"res": {"table_res_list": []}}]) is False
    assert scan_structure_contains_table([{"res": {"table_res_list": [{"pred_html": "<table></table>"}]}}]) is True


def test_borderless_ocr_geometry_only_triggers_structure_hint():
    from cabqp.modules.document_intelligence.ocr.base import OcrLine
    from cabqp.modules.document_intelligence.table import looks_like_table_ocr_lines

    lines = []
    for row, y in enumerate((10, 40, 70, 100)):
        lines.extend([
            OcrLine(str(row), 0.99, (10, y, 25, y + 18), 1, "fake"),
            OcrLine(f"Nguyễn Văn {row}", 0.99, (80, y, 180, y + 18), 1, "fake"),
            OcrLine(f"00{row}123456789", 0.99, (240, y, 340, y + 18), 1, "fake"),
        ])

    assert looks_like_table_ocr_lines(lines) is True


def test_real_digital_pdf_table_is_split_into_people_rows():
    doc = fitz.open()
    page = doc.new_page(width=600, height=300)
    xs = [40, 90, 270, 410, 560]
    ys = [40, 75, 110, 145]
    for x in xs:
        page.draw_line((x, ys[0]), (x, ys[-1]), color=(0, 0, 0), width=1)
    for y in ys:
        page.draw_line((xs[0], y), (xs[-1], y), color=(0, 0, 0), width=1)
    values = [
        ["STT", "Ho ten", "CCCD", "Don vi cong tac"],
        ["1", "Nguyen Van A", "001082946357", "Cuc Ky thuat"],
        ["2", "Tran Thi B", "079203001122", "Trung tam A"],
    ]
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            page.insert_text((xs[c] + 4, ys[r] + 22), value, fontsize=9, fontname="helv")
    content = doc.tobytes()
    doc.close()

    headers, rows, _ = bulk_service._read_rows("people.pdf", content)

    assert "Ho ten" in headers
    assert [(row["Ho ten"], row["CCCD"]) for row in rows] == [
        ("Nguyen Van A", "001082946357"),
        ("Tran Thi B", "079203001122"),
    ]



def test_excel_blank_rows_keep_real_source_row_number():
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Họ và tên", "CCCD"])
    ws.append([])
    ws.append(["Nguyễn Văn A", "001082946357"])
    out = BytesIO(); wb.save(out)

    _, rows, _ = bulk_service._read_rows("people.xlsx", out.getvalue())

    assert len(rows) == 1
    assert rows[0]["_source_row"] == 3
    assert rows[0]["_source_row_index"] == 3


def test_excel_group_heading_does_not_replace_real_field_header():
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Thông tin cá nhân", None, "Công tác", None])
    ws.append(["Họ và tên", "CCCD", "Chức vụ", "Đơn vị công tác"])
    ws.append(["Nguyễn Văn A", "001082946357", "Chuyên viên", "Cục Kỹ thuật"])
    out = BytesIO(); wb.save(out)

    headers, rows, _ = bulk_service._read_rows("people.xlsx", out.getvalue())

    assert headers[:4] == ["Họ và tên", "CCCD", "Chức vụ", "Đơn vị công tác"]
    assert rows[0]["Họ và tên"] == "Nguyễn Văn A"
    assert rows[0]["_source_row"] == 3


def test_mixed_ruled_and_ordinary_pdf_tables_preserve_page_order_and_continuation():
    # Page 1 has enough ruled lines to take the Camelot-lattice branch. Page 2 has
    # fewer horizontal lines, so it takes the pdfplumber branch and omits the header.
    # Returning engine groups instead of page order used to drop page-2 people.
    doc = fitz.open()
    xs = [40, 90, 260, 390, 560]

    page1 = doc.new_page(width=600, height=360)
    ys1 = [35, 70, 105, 140, 175, 210]
    for x in xs: page1.draw_line((x, ys1[0]), (x, ys1[-1]), color=(0,0,0), width=1)
    for y in ys1: page1.draw_line((xs[0], y), (xs[-1], y), color=(0,0,0), width=1)
    page1_rows = [
        ["STT", "Ho ten", "CCCD", "Don vi cong tac"],
        ["1", "Nguyen Van A", "001082946351", "Cuc Ky thuat"],
        ["2", "Tran Thi B", "001082946352", "Cuc Ky thuat"],
        ["3", "Le Van C", "001082946353", "Cuc Ky thuat"],
        ["4", "Pham Thi D", "001082946354", "Cuc Ky thuat"],
    ]
    for r,row in enumerate(page1_rows):
        for c,value in enumerate(row): page1.insert_text((xs[c]+4, ys1[r]+22), value, fontsize=8, fontname="helv")

    page2 = doc.new_page(width=600, height=220)
    ys2 = [35, 70, 105]
    for x in xs: page2.draw_line((x, ys2[0]), (x, ys2[-1]), color=(0,0,0), width=1)
    for y in ys2: page2.draw_line((xs[0], y), (xs[-1], y), color=(0,0,0), width=1)
    page2_rows = [
        ["5", "Do Van E", "001082946355", "Cuc Ky thuat"],
        ["6", "Hoang Thi F", "001082946356", "Cuc Ky thuat"],
    ]
    for r,row in enumerate(page2_rows):
        for c,value in enumerate(row): page2.insert_text((xs[c]+4, ys2[r]+22), value, fontsize=8, fontname="helv")
    content=doc.tobytes(); doc.close()

    from cabqp.modules.document_intelligence.table import extract_pdf_table_results
    native = extract_pdf_table_results(content)
    assert [x["page"] for x in native] == [1,2]

    _, rows, _ = bulk_service._read_rows("people.pdf", content)
    assert [row["Ho ten"] for row in rows] == [
        "Nguyen Van A", "Tran Thi B", "Le Van C", "Pham Thi D", "Do Van E", "Hoang Thi F"
    ]
    assert [row["_source_page"] for row in rows] == [1,1,1,1,2,2]


def test_csv_router_enforces_row_limit(monkeypatch):
    monkeypatch.setattr(
        router,
        "get_settings",
        lambda: SimpleNamespace(document_max_spreadsheet_rows=2, document_max_worksheets=5),
    )
    with pytest.raises(DocumentLimitError) as exc:
        router.probe_tabular("people.csv", b"name,code\nA,1\nB,2\n", "text")
    assert exc.value.code == "SPREADSHEET_ROW_LIMIT_EXCEEDED"


def test_hybrid_table_results_are_sorted_across_native_and_scan_engines(monkeypatch):
    from cabqp.modules.document_intelligence.ocr.base import OcrLine
    from cabqp.modules.document_intelligence.ocr.pipeline import OcrRun
    from cabqp.modules.document_intelligence.router import PageProbe

    doc=fitz.open(); doc.new_page(); p2=doc.new_page(); p2.insert_text((72,72),"digital continuation text "*5)
    content=doc.tobytes(); doc.close()
    decision=RoutingDecision(
        InputKind.PDF_HYBRID,"pdf","test",
        page_route_map=[
            PageProbe(1,"OCR",0,1.0,0,0.0,0.0,0,[]),
            PageProbe(2,"TEXT",100,0.0,1,4.0,0.0,0,[]),
        ],
    )
    settings=SimpleNamespace(
        document_soft_timeout_seconds=60, document_max_ocr_pages=5,
        document_max_image_pixels=40_000_000, document_max_raster_pixels=100_000_000,
        scan_table_detection_enabled=True, document_failed_gate_confidence_cap=0.5,
        document_max_tables=20,
    )
    monkeypatch.setattr(parsers,"route_input",lambda *_: decision)
    monkeypatch.setattr(parsers,"get_settings",lambda: settings)
    monkeypatch.setattr(parsers,"run_ocr",lambda *_,**__: OcrRun(
        [OcrLine("table",0.99,(0,0,100,20),1,"fake")],
        {"gate_result":"PASS","metrics":{}},0.99,
        {"selected_engine":"fake","critical_disagreement":False,"fallback_ran":False},
    ))
    structure=[{"res":{"table_res_list":[{
        "pred_html":"<table><tr><th>Ho ten</th><th>CCCD</th></tr><tr><td>Nguyen Van A</td><td>001082946351</td></tr></table>",
        "table_ocr_pred":{"rec_scores":[0.99,0.99,0.99,0.99]},
    }]}}]
    monkeypatch.setattr(parsers,"looks_like_table_ocr_lines",lambda *_: True)
    monkeypatch.setattr(parsers,"extract_scan_structure",lambda *_,**__: structure)
    monkeypatch.setattr(parsers,"extract_pdf_table_results",lambda *_:[{
        "cells":[["Do Van E","001082946355"]],"page":2,"source":"pdfplumber"
    }])

    result=parsers.parse_document("hybrid.pdf",content)

    assert result.evidence["table_sources"][0]["page"] == 1
    assert result.evidence["table_sources"][1]["page"] == 2
    assert result.tables[0][0] == ["Ho ten","CCCD"]


def test_instruction_sheet_is_ignored_when_a_personnel_sheet_is_detected():
    wb=Workbook()
    guide=wb.active; guide.title="Hướng dẫn"
    guide.append(["Mục", "Nội dung"])
    guide.append(["1", "Nhập dữ liệu tại sheet Danh sách"])
    data=wb.create_sheet("Danh sách")
    data.append(["Họ và tên", "CCCD"])
    data.append(["Nguyễn Văn A", "001082946357"])
    out=BytesIO(); wb.save(out)

    headers,rows,_=bulk_service._read_rows("people.xlsx",out.getvalue())

    assert headers == ["Họ và tên","CCCD"]
    assert len(rows) == 1
    assert rows[0]["_source_sheet"] == "Danh sách"
    assert rows[0]["_source_row"] == 2


def test_pp_structure_config_key_named_table_does_not_fake_detection():
    from cabqp.modules.document_intelligence.table import scan_structure_contains_table

    structure=[{"res":{"model_settings":{"use_table_recognition":False},"table_res_list":[]}}]
    assert scan_structure_contains_table(structure) is False


def test_pdf_continuation_schema_is_not_carried_across_page_gap(monkeypatch):
    first=[
        ["Họ và tên","CCCD"],
        ["Nguyễn Văn A","001082946357"],
    ]
    unrelated=[
        ["Người ký xác nhận","012345678901"],
    ]
    fake=SimpleNamespace(
        tables=[first,unrelated],
        evidence={"table_sources":[{"source":"pdfplumber","page":1},{"source":"pdfplumber","page":3}]},
    )
    monkeypatch.setattr(parsers,"parse_document",lambda *_: fake)

    _,rows,_=bulk_service._read_rows("people.pdf",b"fake")

    assert [row["Họ và tên"] for row in rows] == ["Nguyễn Văn A"]


def test_borderless_pdf_text_layout_is_reconstructed_as_personnel_table(monkeypatch):
    text="""DANH SÁCH NHÂN SỰ
STT
HỌ VÀ TÊN
MÃ CÁ NHÂN
NĂM SINH
CHỨC VỤ
ĐƠN VỊ CÔNG TÁC
1
Lê Hoàng Duy
SYN-CAND-000003
1986
Bác sĩ
Học viện An ninh nhân dân
2
Phạm Gia Huy
SYN-QĐND-000004
1987
Trợ lý
Bệnh viện Quân y 175
Ghi chú cuối trang
"""
    table=bulk_service.extract_unruled_text_table(text)

    assert table[0] == ["STT","HỌ VÀ TÊN","MÃ CÁ NHÂN","NĂM SINH","CHỨC VỤ","ĐƠN VỊ CÔNG TÁC"]
    assert len(table) == 3
    assert bulk_service.tables_have_record_list([table]) is True

    fake=SimpleNamespace(tables=[],text=text,evidence={})
    monkeypatch.setattr(parsers,"parse_document",lambda *_: fake)
    headers,rows,_=bulk_service._read_rows("people.pdf",b"fake")
    mapping,_=bulk_service.infer_mapping(headers,rows)

    assert mapping["HỌ VÀ TÊN"] == "subject_name"
    assert [row["HỌ VÀ TÊN"] for row in rows] == ["Lê Hoàng Duy","Phạm Gia Huy"]


def test_unruled_text_fallback_rejects_ordinary_numbered_document():
    text="""QUYẾT ĐỊNH
Điều 1
Nội dung quyết định
Điều 2
Hiệu lực thi hành
"""
    assert bulk_service.extract_unruled_text_table(text) == []
