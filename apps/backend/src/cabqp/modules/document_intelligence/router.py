from __future__ import annotations

import csv
import math
import re
from dataclasses import asdict, dataclass, field
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook

from cabqp.modules.document_intelligence.file_validation import detect_type
from cabqp.modules.document_intelligence.limits import DocumentLimitError
from cabqp.shared.enums import InputKind
from cabqp.shared.settings import get_settings


def _entropy(text: str) -> float:
    chars=[c for c in text if not c.isspace()]
    if not chars: return 0.0
    freq={}
    for c in chars: freq[c]=freq.get(c,0)+1
    n=len(chars)
    return -sum((v/n)*math.log2(v/n) for v in freq.values())


@dataclass(slots=True)
class PageProbe:
    page: int
    route: str
    chars: int
    image_coverage: float
    embedded_fonts: int
    char_entropy: float
    replacement_ratio: float
    rotation: int
    reasons: list[str]=field(default_factory=list)


@dataclass(slots=True)
class RoutingDecision:
    kind: InputKind
    detected_type: str
    reason: str
    page_route_map: list[PageProbe]=field(default_factory=list)
    tabular_profile: dict=field(default_factory=dict)

    def to_dict(self) -> dict:
        return {
            "kind": self.kind.value,
            "detected_type": self.detected_type,
            "reason": self.reason,
            "page_route_map": [asdict(p) for p in self.page_route_map],
            "tabular_profile": self.tabular_profile,
        }


def _sample_indices(n: int) -> list[int]:
    if n <= 3: return list(range(n))
    return sorted(set([0, n//2, n-1]))


def probe_pdf(content: bytes) -> RoutingDecision:
    try:
        import fitz
        doc=fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        # A file that only *starts* with "%PDF-" passes magic-byte validation, so a
        # truncated or corrupt upload reaches here. Treating it as a scan was not a
        # safe fallback: every scan path rasterizes through this same library, so the
        # document could never be read — it just failed later, with OCR blamed for it.
        # Refuse it as a document instead, which the upload route answers 422 to and
        # the worker turns into an operator-visible NEED_REVIEW.
        raise DocumentLimitError("INVALID_PDF_DOCUMENT", reason=type(exc).__name__) from exc
    settings=get_settings()
    if len(doc) > settings.document_max_pdf_pages:
        pages=len(doc); doc.close()
        raise DocumentLimitError("PDF_PAGE_LIMIT_EXCEEDED", pages=pages, limit=settings.document_max_pdf_pages)
    probes=[]
    routes=[]
    for i in range(len(doc)):
        page=doc[i]
        text=page.get_text("text") or ""
        chars=len(text.strip())
        repl=text.count("�")/max(1,len(text))
        ent=_entropy(text)
        try: fonts=len(page.get_fonts(full=True))
        except Exception: fonts=0
        coverage=0.0
        try:
            area=max(1.0, float(page.rect.width*page.rect.height))
            boxes=[]
            for info in page.get_image_info(xrefs=True):
                b=info.get("bbox")
                if b:
                    boxes.append(max(0.0,(b[2]-b[0])*(b[3]-b[1])))
            coverage=min(1.0,sum(boxes)/area)
        except Exception:
            pass
        reasons=[]
        scan_score=0
        if chars < 50: scan_score+=1; reasons.append("chars<50")
        if coverage > 0.75: scan_score+=2; reasons.append("bitmap_coverage>0.75")
        if fonts == 0: scan_score+=1; reasons.append("no_embedded_fonts")
        if chars >= 50 and ent < 3.5: scan_score+=2; reasons.append("low_char_entropy")
        if repl > 0.02: scan_score+=2; reasons.append("replacement_ratio>0.02")
        route="OCR" if scan_score >= 2 else "TEXT"
        # text-as-paths often looks digital but has no extractable chars or raster image.
        if chars < 50 and coverage < 0.1 and fonts == 0:
            route="OCR"; reasons.append("possible_text_as_paths")
        probes.append(PageProbe(i+1,route,chars,coverage,fonts,ent,repl,int(page.rotation or 0),reasons))
        routes.append(route)
    doc.close()
    if not probes:
        return RoutingDecision(InputKind.PDF_SCAN,"pdf","empty_pdf")
    if all(x=="TEXT" for x in routes): kind=InputKind.PDF_TEXT
    elif all(x=="OCR" for x in routes): kind=InputKind.PDF_SCAN
    else: kind=InputKind.PDF_HYBRID
    return RoutingDecision(kind,"pdf",f"pdf_per_page_probe:{kind.value}",probes)


def _is_stringish(v) -> bool:
    return isinstance(v,str) and bool(v.strip())


def _profile_matrix(rows: list[list]) -> tuple[InputKind,dict]:
    clean=[]
    for r in rows:
        vals=list(r)
        while vals and vals[-1] in (None,""): vals.pop()
        if any(v not in (None,"") for v in vals): clean.append(vals)
    if not clean:
        return InputKind.KEY_VALUE_SHEET,{"reason":"empty"}
    header_idx=None
    for i,row in enumerate(clean[:20]):
        non=[v for v in row if v not in (None,"")]
        if len(non)>=3 and sum(_is_stringish(v) for v in non)/len(non)>=0.8:
            header_idx=i; break
    n_cols=max((len(r) for r in clean),default=0)
    first_col_labels=sum(
        isinstance(r[0],str) and r[0].strip().endswith(":")
        for r in clean if r
    )/max(1,len(clean))
    if n_cols <= 2 or first_col_labels >= 0.3:
        return InputKind.KEY_VALUE_SHEET,{"header_row":header_idx,"n_cols":n_cols,"label_ratio":first_col_labels,"reason":"key_value_antisignal"}
    if header_idx is None:
        return InputKind.KEY_VALUE_SHEET,{"header_row":None,"n_cols":n_cols,"reason":"no_header"}
    header=clean[header_idx]
    expected=len(header)
    data=clean[header_idx+1:]
    rectangular=sum(abs(len(r)-expected)<=1 for r in data)/max(1,len(data))
    # unique key signal: CCCD/CMND-like column or near-unique column
    key_signal=False
    for c in range(expected):
        vals=[str(r[c]).strip() for r in data if c < len(r) and r[c] not in (None,"")]
        if len(vals)>=2:
            near_unique=len(set(vals))/len(vals)>=0.9
            id_like=sum(bool(re.fullmatch(r"\d{9}|\d{12}",v)) for v in vals)/len(vals)>=0.6
            key_signal=key_signal or near_unique or id_like
    kind=InputKind.TABULAR_LIST if len(data)>=2 and rectangular>=0.7 else InputKind.KEY_VALUE_SHEET
    return kind,{"header_row":header_idx,"n_cols":expected,"data_rows":len(data),"rectangularity":rectangular,"key_signal":key_signal,"reason":"rectangular_list" if kind==InputKind.TABULAR_LIST else "non_rectangular"}


def probe_tabular(file_name: str, content: bytes, detected: str) -> RoutingDecision:
    rows=[]
    if detected == "xlsx":
        wb=load_workbook(BytesIO(content),read_only=True,data_only=True)
        settings=get_settings()
        if len(wb.worksheets) > settings.document_max_worksheets:
            raise DocumentLimitError("WORKSHEET_LIMIT_EXCEEDED", worksheets=len(wb.worksheets), limit=settings.document_max_worksheets)
        total_rows=0
        for ws in wb.worksheets:
            sheet_rows=[list(r) for r in ws.iter_rows(values_only=True)]
            total_rows += len(sheet_rows)
            if total_rows > settings.document_max_spreadsheet_rows:
                raise DocumentLimitError("SPREADSHEET_ROW_LIMIT_EXCEEDED", rows=total_rows, limit=settings.document_max_spreadsheet_rows)
            rows.extend(sheet_rows)
    elif detected == "xls":
        import xlrd
        book=xlrd.open_workbook(file_contents=content)
        settings=get_settings()
        if book.nsheets > settings.document_max_worksheets:
            raise DocumentLimitError("WORKSHEET_LIMIT_EXCEEDED", worksheets=book.nsheets, limit=settings.document_max_worksheets)
        total_rows=0
        for sh in book.sheets():
            total_rows += sh.nrows
            if total_rows > settings.document_max_spreadsheet_rows:
                raise DocumentLimitError("SPREADSHEET_ROW_LIMIT_EXCEEDED", rows=total_rows, limit=settings.document_max_spreadsheet_rows)
            rows.extend([sh.row_values(i) for i in range(sh.nrows)])
    else:
        text=content.decode("utf-8-sig",errors="replace")
        try: dialect=csv.Sniffer().sniff(text[:4096],delimiters=",;\t|")
        except csv.Error: dialect=csv.excel
        limit=get_settings().document_max_spreadsheet_rows
        for row_no,row in enumerate(csv.reader(StringIO(text),dialect),start=1):
            if row_no > limit:
                raise DocumentLimitError("SPREADSHEET_ROW_LIMIT_EXCEEDED", rows=row_no, limit=limit)
            rows.append(row)
    kind,profile=_profile_matrix(rows)
    return RoutingDecision(kind,detected,f"tabular_probe:{kind.value}",tabular_profile=profile)


def route_input(file_name: str, content: bytes) -> RoutingDecision:
    detected=detect_type(content)
    ext=Path(file_name).suffix.lower()
    # Hard rule: spreadsheet/CSV always receives the tabular probe.
    if ext in {".xlsx",".xls",".csv"} and detected in {"xlsx","xls","text"}:
        return probe_tabular(file_name,content,detected)
    if detected == "pdf": return probe_pdf(content)
    if detected == "docx": return RoutingDecision(InputKind.DOCX,detected,"docx_structured_first")
    if detected in {"png","jpeg"}: return RoutingDecision(InputKind.IMAGE,detected,"image_requires_ocr")
    if detected == "text": return RoutingDecision(InputKind.TEXT,detected,"plain_text")
    return RoutingDecision(InputKind.IMAGE,detected,"unknown_binary_requires_review")
